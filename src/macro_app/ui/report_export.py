"""Xuất scorecard ra PDF và Word: gauge, điểm, vì sao, độ tin cậy, bảng chỉ số theo nhóm.

Không phụ thuộc Streamlit: trang Scorecard gom số vào `ReportData` rồi gọi `to_pdf` / `to_docx`.
PDF dùng font DejaVu Sans đi kèm matplotlib (đủ dấu tiếng Việt, không phụ thuộc font máy chủ).
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd

from macro_app import fmt
from macro_app.charts import theme as t

SCORE_NAME = {2: "Rất tốt", 1: "Tốt", 0: "Trung tính", -1: "Xấu", -2: "Rất xấu"}
SCORE_STATUS = {2: "green_strong", 1: "green", 0: "yellow", -1: "orange", -2: "red"}
RATING_STATUS = {
    "Rất thuận lợi": "green_strong",
    "Thuận lợi": "green",
    "Trung tính": "yellow",
    "Bất lợi": "orange",
    "Rất bất lợi": "red",
}


@dataclass
class Row:
    name: str
    updated: str
    value: str
    level: str
    status: str  # khóa STATUS_COLORS
    compare: str
    reference: str
    weight: str
    contribution: str


@dataclass
class Group:
    name: str
    level: str
    status: str
    contribution: str
    rows: list[Row] = field(default_factory=list)


@dataclass
class ReportData:
    profile: str
    segment: str
    asof: str
    compare_label: str
    score: float
    rating: str
    score_note: str
    saved_note: str
    coverage: str
    stale: list[str]
    sensitivity: str
    drivers: list[tuple[str, float]]
    moves: str
    groups: list[Group]
    bands: list[dict]
    footnote: str
    # mỗi mục: tông (bad, good hoặc rỗng), tên chỉ số, câu hàm ý
    implications: list[tuple[str, str, str]] = field(default_factory=list)
    meaning: str = ""  # bất lợi / thuận lợi nghĩa là tăng hay giảm, trong bao lâu


def _esc(text: str) -> str:
    """Chữ đưa vào Paragraph của reportlab (hiểu thẻ kiểu HTML)."""
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _bucket(x) -> int | None:
    if x is None or pd.isna(x):
        return None
    return int(max(-2, min(2, np.floor(float(x) + 0.5))))


# --- gauge (ảnh PNG, dùng chung cho PDF và Word) ---


def gauge_png(score: float, bands: list[dict], width_in: float = 3.2) -> bytes:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Polygon, Wedge

    lo, hi = -1.0, 1.2

    def ang(v: float) -> float:  # độ, 180 = trái
        return 180 * (1 - (np.clip(v, lo, hi) - lo) / (hi - lo))

    fig, ax = plt.subplots(figsize=(width_in, width_in * 0.62), dpi=200)
    tops = [hi, *[b["min"] for b in bands[:-1]]]
    for band, top in zip(bands, tops, strict=True):
        start, end = max(band["min"], lo), min(top, hi)
        if end > start:
            ax.add_patch(
                Wedge(
                    (0, 0),
                    1.0,
                    ang(end) + 0.4,
                    ang(start) - 0.4,
                    width=0.26,
                    color=t.RATING_SCALE[band["label"]],
                    lw=0,
                )
            )
    for b in bands:
        if lo < b["min"] < hi:
            a = np.radians(ang(b["min"]))
            ax.text(
                1.14 * np.cos(a),
                1.14 * np.sin(a),
                fmt.number(b["min"], 2),
                ha="center",
                va="center",
                fontsize=6.5,
                color=t.MUTED,
            )
    ax.text(-1.12, -0.08, "≤ −1", ha="center", fontsize=6.5, color=t.MUTED)
    ax.text(1.12, -0.08, "≥ 1,2", ha="center", fontsize=6.5, color=t.MUTED)
    if pd.notna(score):
        a = np.radians(ang(score))
        tip = (0.70 * np.cos(a), 0.70 * np.sin(a))
        side = a + np.pi / 2
        left = (0.045 * np.cos(side), 0.045 * np.sin(side))
        right = (-left[0], -left[1])
        tail = (-0.13 * np.cos(a), -0.13 * np.sin(a))
        ax.add_patch(Polygon([tip, left, tail, right], closed=True, color=t.HEADING, lw=0))
    ax.add_patch(Wedge((0, 0), 0.085, 0, 360, color=t.HEADING, lw=0))
    ax.add_patch(Wedge((0, 0), 0.035, 0, 360, color="white", lw=0))
    ax.set_xlim(-1.3, 1.3)
    ax.set_ylim(-0.2, 1.2)
    ax.set_aspect("equal")
    ax.axis("off")
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight", pad_inches=0.02, transparent=True)
    plt.close(fig)
    return buf.getvalue()


# --- PDF ---


def _fonts() -> tuple[str, str]:
    import matplotlib
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    base = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
    if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVu", str(base / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(base / "DejaVuSans-Bold.ttf")))
    return "DejaVu", "DejaVu-Bold"


def _pdf_implications(d: ReportData, p, gap) -> list:
    """Khối "Hàm ý" của PDF: mỗi chỉ số một dòng, bất lợi ▼ / thuận lợi ▲."""
    if not d.implications:
        return []
    marks = {"bad": ("▼", t.STATUS_COLORS["orange"][1]), "good": ("▲", t.STATUS_COLORS["green"][1])}
    out = [p("AI comment · tác động lên Việt Nam", size=10, fnt=_fonts()[1], color=t.HEADING)]
    for tone, name, text in d.implications:
        mark, col = marks.get(tone, ("•", t.MUTED))
        out.append(
            p(f'<font color="{col}">{mark}</font> <b>{_esc(name)}</b>: {_esc(text)}', size=8)
        )
    return [*out, gap]


def to_pdf(d: ReportData) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, bold = _fonts()
    c = colors.HexColor

    def style(size=8.5, color=t.TEXT, fnt=None, lead=None, align=0):
        return ParagraphStyle(
            "s",
            fontName=fnt or font,
            fontSize=size,
            leading=lead or size * 1.3,
            textColor=c(color),
            alignment=align,
        )

    def p(text, **kw):
        return Paragraph(str(text), style(**kw))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=8 * mm,
        bottomMargin=7 * mm,
        title=f"Scorecard {d.profile} - {d.segment}",
    )
    story = [
        p(f"Scorecard · {d.profile} · {d.segment}", size=15, fnt=bold, color=t.HEADING),
        p(f"Ngày {d.asof} · so với {d.compare_label} · {d.saved_note}", size=8, color=t.MUTED),
        p(_esc(d.meaning), size=8.5) if d.meaning else Spacer(1, 1),
        Spacer(1, 4 * mm),
    ]
    _, fg = t.STATUS_COLORS[RATING_STATUS.get(d.rating, "none")]
    score_txt = fmt.signed(d.score, 2) if pd.notna(d.score) else "—"
    hero_text = [
        p(f'<font name="{bold}" size="26" color="{t.TEXT}">{score_txt}</font>', lead=30),
        p(f'<font color="{fg}" name="{bold}">{d.rating}</font>', size=11),
        Spacer(1, 2 * mm),
        p(d.score_note, size=8, color=t.MUTED),
        Spacer(1, 3 * mm),
        p(f"<b>Độ phủ trọng số:</b> {d.coverage}", size=8.5),
        p("<b>Chỉ số số cũ (không tính):</b> " + (", ".join(d.stale) or "không có"), size=8.5),
        p(d.sensitivity, size=8.5) if d.sensitivity else Spacer(1, 1),
    ]
    why_rows = [[p("Nhóm", fnt=bold, size=8), p("Đóng góp", fnt=bold, size=8, align=2)]]
    for name, val in d.drivers:
        col = t.STATUS_COLORS["green"][1] if val >= 0 else t.STATUS_COLORS["orange"][1]
        why_rows.append(
            [
                p(name, size=8.5),
                p(f'<font color="{col}">{fmt.signed(val, 2)}</font>', size=8.5, align=2),
            ]
        )
    why = Table(why_rows, colWidths=[62 * mm, 22 * mm])
    why.setStyle(
        TableStyle(
            [
                ("LINEBELOW", (0, 0), (-1, 0), 0.8, c(t.HEADING)),
                ("LINEBELOW", (0, 1), (-1, -1), 0.3, c(t.BORDER)),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    why_box = [
        p("Vì sao điểm thế này", size=10, fnt=bold, color=t.HEADING),
        Spacer(1, 1.5 * mm),
        why,
        Spacer(1, 2 * mm),
        p(f"<b>Đổi so {d.compare_label}:</b> {d.moves}", size=8),
    ]
    gauge = Image(io.BytesIO(gauge_png(d.score, d.bands)), width=62 * mm, height=38 * mm)
    hero = Table([[gauge, hero_text, why_box]], colWidths=[76 * mm, 80 * mm, 115 * mm])
    hero.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOX", (0, 0), (-1, -1), 0.5, c(t.BORDER)),
                ("BACKGROUND", (0, 0), (-1, -1), c("#F7F9FA")),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story += [
        hero,
        Spacer(1, 2.5 * mm),
        *_pdf_implications(d, p, Spacer(1, 2.5 * mm)),
        p("Chi tiết chỉ số", size=11, fnt=bold, color=t.HEADING),
        Spacer(1, 1.5 * mm),
    ]
    head = [
        "Chỉ số",
        "Cập nhật",
        "Giá trị dùng chấm",
        "Mức",
        f"So {d.compare_label}",
        "Thông tin tham khảo",
        "Trọng số",
        "Đóng góp",
    ]
    data = [[p(h, fnt=bold, size=7.5, color=t.MUTED) for h in head]]
    styles = [
        ("LINEBELOW", (0, 0), (-1, 0), 1, c(t.HEADING)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (2, 0), (2, -1), "RIGHT"),
        ("ALIGN", (6, 0), (7, -1), "RIGHT"),
        ("TOPPADDING", (0, 0), (-1, -1), 1.6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6),
    ]
    for g in d.groups:
        i = len(data)
        gbg, gfg = t.STATUS_COLORS.get(g.status, t.STATUS_COLORS["none"])
        data.append(
            [
                p(g.name, fnt=bold, size=8.5),
                "",
                "",
                p(g.level, size=8, fnt=bold, color=gfg),
                "",
                "",
                "",
                p(g.contribution, fnt=bold, size=8.5, align=2),
            ]
        )
        styles += [
            ("BACKGROUND", (0, i), (-1, i), c("#F1F4F6")),
            ("BACKGROUND", (3, i), (3, i), c(gbg)),
            ("LINEBELOW", (0, i), (-1, i), 0.4, c(t.BORDER)),
        ]
        for r in g.rows:
            j = len(data)
            rbg, rfg = t.STATUS_COLORS.get(r.status, t.STATUS_COLORS["none"])
            data.append(
                [
                    p(r.name, size=8),
                    p(r.updated, size=7.5, color=t.MUTED),
                    p(r.value, size=8.5, fnt=bold, align=2),
                    p(r.level, size=8, fnt=bold, color=rfg),
                    p(r.compare, size=8),
                    p(r.reference, size=7.5),
                    p(r.weight, size=8, align=2),
                    p(r.contribution, size=8, align=2),
                ]
            )
            styles += [
                ("BACKGROUND", (3, j), (3, j), c(rbg)),
                ("LINEBELOW", (0, j), (-1, j), 0.25, c(t.BORDER)),
            ]
    table = Table(
        data,
        colWidths=[64 * mm, 22 * mm, 30 * mm, 30 * mm, 26 * mm, 52 * mm, 20 * mm, 22 * mm],
        repeatRows=1,
    )
    table.setStyle(TableStyle(styles))
    story += [table, Spacer(1, 4 * mm), p(d.footnote, size=7, color=t.MUTED)]
    doc.build(story)
    return buf.getvalue()


# --- Word ---


def _shade(cell, hex_color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tc.append(shd)


def _light_borders(table) -> None:
    """Viền bảng xám nhạt như bản PDF, thay đường kẻ đen mặc định của Word."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil" if edge in ("left", "right", "insideV") else "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), t.BORDER.lstrip("#"))
        borders.append(el)
    table._tbl.tblPr.append(borders)


def _run(par, text: str, *, size=9, bold=False, color=t.TEXT):
    from docx.shared import Pt, RGBColor

    run = par.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor.from_string(color.lstrip("#").upper())
    return run


def to_docx(d: ReportData) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.3))
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(9)

    _run(
        doc.add_paragraph(),
        f"Scorecard · {d.profile} · {d.segment}",
        size=16,
        bold=True,
        color=t.HEADING,
    )
    _run(
        doc.add_paragraph(),
        f"Ngày {d.asof} · so với {d.compare_label} · {d.saved_note}",
        size=8,
        color=t.MUTED,
    )

    if d.meaning:
        _run(doc.add_paragraph(), d.meaning, size=9)
    _docx_hero(doc, d)
    if d.implications:
        _run(
            doc.add_paragraph(),
            "AI comment · tác động lên Việt Nam",
            size=10.5,
            bold=True,
            color=t.HEADING,
        )
        for tone, name, text in d.implications:
            par = doc.add_paragraph()
            par.paragraph_format.space_after = Pt(1)
            mark = {"bad": "▼ ", "good": "▲ "}.get(tone, "• ")
            col = {"bad": t.STATUS_COLORS["orange"][1], "good": t.STATUS_COLORS["green"][1]}.get(
                tone, t.MUTED
            )
            _run(par, mark, size=8.5, color=col)
            _run(par, f"{name}: ", size=8.5, bold=True)
            _run(par, text, size=8.5)
    _docx_table(doc, d)
    _run(doc.add_paragraph(), d.footnote, size=7.5, color=t.MUTED)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _docx_hero(doc, d: ReportData) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Cm

    hero = doc.add_table(rows=1, cols=3)
    hero.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, mid, right = hero.rows[0].cells
    for cell in (left, mid, right):
        _shade(cell, "#F7F9FA")
    left.paragraphs[0].add_run().add_picture(io.BytesIO(gauge_png(d.score, d.bands)), width=Cm(7.2))
    _, fg = t.STATUS_COLORS[RATING_STATUS.get(d.rating, "none")]
    par = mid.paragraphs[0]
    _run(par, fmt.signed(d.score, 2) if pd.notna(d.score) else "—", size=26, bold=True)
    _run(par, f"   {d.rating}", size=12, bold=True, color=fg)
    _run(mid.add_paragraph(), d.score_note, size=8, color=t.MUTED)
    _run(mid.add_paragraph(), f"Độ phủ trọng số: {d.coverage}", size=9)
    _run(
        mid.add_paragraph(),
        "Chỉ số số cũ (không tính): " + (", ".join(d.stale) or "không có"),
        size=9,
    )
    if d.sensitivity:
        _run(mid.add_paragraph(), d.sensitivity.replace("<b>", "").replace("</b>", ""), size=9)
    _run(right.paragraphs[0], "Vì sao điểm thế này", size=11, bold=True, color=t.HEADING)
    for name, val in d.drivers:
        col = t.STATUS_COLORS["green"][1] if val >= 0 else t.STATUS_COLORS["orange"][1]
        line = right.add_paragraph()
        _run(line, f"{name}: ", size=9)
        _run(line, fmt.signed(val, 2), size=9, bold=True, color=col)
    _run(right.add_paragraph(), f"Đổi so {d.compare_label}: {d.moves}", size=8, color=t.MUTED)
    left.width, mid.width, right.width = Cm(7.8), Cm(8.5), Cm(10.5)


def _docx_table(doc, d: ReportData) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    _run(doc.add_paragraph(), "Chi tiết chỉ số", size=12, bold=True, color=t.HEADING)
    head = [
        "Chỉ số",
        "Cập nhật",
        "Giá trị dùng chấm",
        "Mức",
        f"So {d.compare_label}",
        "Thông tin tham khảo",
        "Trọng số",
        "Đóng góp",
    ]
    table = doc.add_table(rows=1, cols=len(head))
    table.style = "Table Grid"
    _light_borders(table)
    for cell, text in zip(table.rows[0].cells, head, strict=True):
        _run(cell.paragraphs[0], text, size=8, bold=True, color=t.MUTED)
        _shade(cell, "#E8EDF1")
    for g in d.groups:
        cells = table.add_row().cells
        for cell in cells:
            _shade(cell, "#F1F4F6")
        _run(cells[0].paragraphs[0], g.name, size=9, bold=True)
        gbg, gfg = t.STATUS_COLORS.get(g.status, t.STATUS_COLORS["none"])
        _shade(cells[3], gbg)
        _run(cells[3].paragraphs[0], g.level, size=8.5, bold=True, color=gfg)
        _run(cells[7].paragraphs[0], g.contribution, size=9, bold=True)
        for r in g.rows:
            cells = table.add_row().cells
            values = [
                r.name,
                r.updated,
                r.value,
                r.level,
                r.compare,
                r.reference,
                r.weight,
                r.contribution,
            ]
            for k, (cell, text) in enumerate(zip(cells, values, strict=True)):
                bg, fg = t.STATUS_COLORS.get(r.status, t.STATUS_COLORS["none"])
                if k == 3:
                    _shade(cell, bg)
                _run(
                    cell.paragraphs[0],
                    text,
                    size=8.5,
                    bold=k in (2, 3),
                    color=fg if k == 3 else t.TEXT,
                )
                if k in (2, 6, 7):
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    widths = [Cm(7.2), Cm(2.4), Cm(3.2), Cm(3.0), Cm(2.8), Cm(5.4), Cm(1.9), Cm(2.0)]
    for row in table.rows:
        for cell, w in zip(row.cells, widths, strict=True):
            cell.width = w


# --- gom số từ trang Scorecard ---

REF_KIND = {"target": "Mục tiêu", "market": "Kỳ vọng"}


def _value(rec: pd.Series, decimals: int) -> str:
    v, unit = rec.get("measured"), rec.get("measured_unit") or ""
    if pd.isna(v):
        return "—"
    if rec.get("measure_text"):
        decimals = max(decimals, 1)
    return f"{fmt.number(v, decimals)}{'%' if unit == '%' else ' ' + unit}"


def _row(rec: pd.Series, ind, lat: pd.Series | None, prev: dict) -> Row:
    info, stale = bool(rec.get("info")), bool(rec.get("stale"))
    if info:
        level, status = "Tham khảo", "none"
    elif stale:
        level, status = "Không tính (số cũ)", "none"
    else:
        level = str(rec.get("label") or "chưa chấm")
        status = rec["status"] if rec["status"] in SCORE_STATUS.values() else "none"
    before = prev.get(rec["code"])
    compare = ""
    if not info and pd.notna(rec.get("score")) and before is not None and pd.notna(before):
        diff = float(rec["score"]) - float(before)
        compare = (
            "=" if abs(diff) < 1e-9 else f"{'▲' if diff > 0 else '▼'} {fmt.signed(diff, 0)} bậc"
        )
    ref = ""
    if lat is not None and pd.notna(rec.get("fwd_value")):
        v = float(rec["fwd_value"])
        unit = lat["unit"]
        kind = REF_KIND.get(str(rec.get("fwd_kind")), "Tham khảo")
        if kind == "Mục tiêu" and unit not in ("%", "điểm %"):
            kind = "Kế hoạch"
        ref = f"{fmt.number(v, 0 if v.is_integer() else ind.decimals)}{'%' if unit == '%' else ' ' + unit} · {kind}"
    counted = not info and not stale and pd.notna(rec.get("score"))
    return Row(
        name=ind.name,
        updated=fmt.period(lat["period"], lat["frequency"]) if lat is not None else "",
        value=_value(rec, ind.decimals),
        level=level,
        status=status,
        compare=compare,
        reference=ref,
        weight=f"{fmt.number(rec.get('weight_used', 0) * 100, 1)}%" if counted else "—",
        contribution=fmt.signed(rec["contribution"], 3) if counted else "—",
    )


def build_report(  # noqa: PLR0913
    *,
    profile: str,
    card: dict,
    table: pd.DataFrame,
    total: dict,
    latest: pd.DataFrame,
    catalog: dict,
    months: int,
    compare_label: str,
    saved_note: str,
    bands: list[dict],
    implications: list[tuple[str, str, str]] = (),
) -> ReportData:
    lat = latest.set_index("code") if not latest.empty else pd.DataFrame()
    cmp = total["compare"][months]
    prev_rows, prev_pillars = cmp.get("rows") or {}, cmp.get("pillars") or {}
    groups = []
    for p in total.get("pillars", []):
        b = _bucket(p["level"])
        part = table[table["pillar"] == p["pillar"]]
        groups.append(
            Group(
                name=p["pillar"],
                level=SCORE_NAME.get(b, "Không có số"),
                status=SCORE_STATUS.get(b, "none"),
                contribution=fmt.signed(p["contribution"], 3)
                if pd.notna(p["contribution"])
                else "—",
                rows=[
                    _row(
                        rec,
                        catalog[rec["code"]],
                        lat.loc[rec["code"]] if rec["code"] in lat.index else None,
                        prev_rows,
                    )
                    for _, rec in part.iterrows()
                ],
            )
        )
    score = total["score"]
    prev = cmp.get("total", np.nan)
    when = pd.Timestamp(cmp["date"]).strftime("%m/%Y") if cmp.get("date") is not None else ""
    score_note = (
        f"{fmt.signed(score - prev, 2)} so với {compare_label} ({when}: {fmt.signed(prev, 2)})"
        if pd.notna(score) and pd.notna(prev)
        else f"Chưa có điểm {compare_label} để so."
    )
    names = {c: i.name for c, i in catalog.items()}
    stale = (
        table[table["stale"].astype(bool) & ~table["info"].astype(bool)]
        if not table.empty
        else table
    )
    sens = total.get("sensitivity") or {}
    sensitivity = (
        f"<b>Nếu số cũ quay lại:</b> giữ mức cuối đã biết → {fmt.signed(sens['keep'], 2)}; "
        f"biên độ {fmt.signed(sens['worst'], 2)} … {fmt.signed(sens['best'], 2)}"
        if sens
        else ""
    )
    drivers = sorted(
        (
            (p["pillar"], p["contribution"])
            for p in total.get("pillars", [])
            if pd.notna(p["contribution"])
        ),
        key=lambda x: x[1],
    )
    moves = " · ".join(
        f"{n} {fmt.signed(v - prev_pillars[n], 2)}"
        for n, v in drivers
        if n in prev_pillars and abs(v - prev_pillars[n]) > 0.005
    ) or ("không đổi" if prev_pillars else f"chưa có số {compare_label}")
    return ReportData(
        profile=profile,
        implications=list(implications),
        meaning=card.get("meaning", ""),
        segment=card["name"],
        asof=pd.Timestamp.now().strftime("%d/%m/%Y"),
        compare_label=compare_label,
        score=score,
        rating=total["rating"],
        score_note=score_note,
        saved_note=saved_note,
        coverage=f"{fmt.number(total.get('coverage', 0) * 100, 0)}%",
        stale=[names.get(c, c) for c in stale["code"]] if len(stale) else [],
        sensitivity=sensitivity,
        drivers=drivers,
        moves=moves,
        groups=groups,
        bands=bands,
        footnote=(
            "Điểm = tổng có trọng số của mức từng chỉ số (+2 đến −2), chia cho độ phủ; tỷ trọng mặc định "
            "chia đều theo nhóm. Chỉ số thiếu hoặc số cũ không tính. Ngưỡng và thang điểm đang chờ "
            "bộ phận QTRR duyệt. Xếp hạng: ≥ 0,8 Rất thuận lợi · ≥ 0,45 Thuận lợi · ≥ 0,05 Trung tính · "
            "≥ −0,2 Bất lợi · thấp hơn Rất bất lợi."
        ),
    )

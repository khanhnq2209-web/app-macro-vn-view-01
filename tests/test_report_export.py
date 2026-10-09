"""Xuất báo cáo PDF/Word: file mở được, 1 trang, đủ nhóm và chỉ số, đúng điểm."""

import io

import pandas as pd

from macro_app.metrics.scorecard import load_settings
from macro_app.ui import report_export as rx

BANDS = load_settings()["rating_bands"]


def sample() -> rx.ReportData:
    rows = [
        rx.Row(
            "LS huy động 12 tháng",
            "09/10/2026",
            "5,90%",
            "Trung tính",
            "yellow",
            "=",
            "",
            "50,0%",
            "±0,000",
        ),
        rx.Row(
            "CPI so cùng kỳ",
            "T8/2026",
            "4,89%",
            "Xấu",
            "orange",
            "▼ −1 bậc",
            "4,5% · Mục tiêu",
            "50,0%",
            "−0,500",
        ),
    ]
    return rx.ReportData(
        profile="Bộ thử",
        segment="Nhà ở",
        asof="09/10/2026",
        compare_label="tháng trước",
        score=-0.5,
        rating="Rất bất lợi",
        score_note="−0,10 so với tháng trước",
        saved_note="bản gốc",
        coverage="100%",
        stale=[],
        sensitivity="",
        drivers=[("Giá", -0.5)],
        moves="Giá −0,10",
        groups=[rx.Group("Giá", "Xấu", "orange", "−0,500", rows)],
        bands=BANDS,
        footnote="Ghi chú",
    )


def test_pdf_is_valid_one_page_with_vietnamese_text():
    import pypdfium2 as pdfium

    data = rx.to_pdf(sample())
    assert data[:5] == b"%PDF-"
    pdf = pdfium.PdfDocument(data)
    assert len(pdf) == 1
    text = pdf[0].get_textpage().get_text_range()
    assert "Scorecard · Bộ thử · Nhà ở" in text and "CPI so cùng kỳ" in text and "−0,50" in text


def test_docx_opens_and_lists_groups_and_rows():
    from docx import Document

    doc = Document(io.BytesIO(rx.to_docx(sample())))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert "Giá" in cells and "CPI so cùng kỳ" in cells and "Xấu" in cells
    assert any("Scorecard · Bộ thử" in p.text for p in doc.paragraphs)


def test_gauge_png_without_score():
    png = rx.gauge_png(float("nan"), BANDS)
    assert png[:8] == b"\x89PNG\r\n\x1a\n"


def test_build_report_from_scorecard_numbers():
    from macro_app.config import load_catalog
    from macro_app.metrics import scorecard as sc

    cat = {i.code: i for i in load_catalog()}
    row = {
        "code": "cpi_yoy",
        "pillar": "Giá",
        "method": "absolute",
        "cuts": [2.5, 3.5, 4.5, 5.5],
        "side": "above",
    }
    card = {"name": "Nhà ở", "rows": [row]}
    idx = pd.date_range("2020-01-31", periods=12, freq="ME")
    frames = {"cpi_yoy": pd.DataFrame({"value": [2.0] * 11 + [6.0], "source": "t"}, index=idx)}
    ctx = {
        "settings": load_settings(),
        "asof": pd.Timestamp("2020-12-31"),
        "min_points": {},
        "targets": {},
        "target_series": {},
    }
    table, total, _ = sc.evaluate_card(card, frames, cat, ctx)
    rep = rx.build_report(
        profile="Bộ",
        card=card,
        table=table,
        total=total,
        latest=pd.DataFrame(),
        catalog=cat,
        months=1,
        compare_label="tháng trước",
        saved_note="x",
        bands=BANDS,
    )
    assert rep.score == -2.0 and rep.rating == "Rất bất lợi"
    assert rep.groups[0].rows[0].compare == "▼ −4 bậc"  # Rất tốt → Rất xấu
    assert rep.groups[0].rows[0].level == "Rất xấu"
